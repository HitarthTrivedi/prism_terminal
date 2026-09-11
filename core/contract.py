"""What a step must PRODUCE — and what Prism does when it doesn't.

The problem this exists to solve
────────────────────────────────
Prism used to describe a step only in prose. The planner wrote "Your ONLY
task is: act as a senior technical writer … Deliverable Spec: … Non-Goals:
do NOT generate an actual .docx file", the engine appended its own rules,
and whatever came back was recorded as that step's result. Nothing ever
asked the one question a person actually cares about:

    the person asked for a document — is there a document?

So a run could end with every step green and no file, no picture and no
video anywhere on the disk. It did, repeatedly: a "make DOCX document"
run that produced 9,521 characters of Word-ready text in a chat window; a
reel run whose image step replied "I can't complete the requested PNG
generation" and was recorded as completed; a Claude step that DID build the
DOCX, which Prism then failed to collect and never mentioned again.

The fix is to stop treating the deliverable as something described in prose
and start treating it as a CONTRACT the step is checked against:

    kind      what has to exist when the step ends
    ────      ───────────────────────────────────────────────────────────
    text      a reply in the chat
    file      a document/deck/sheet/archive, downloaded and saved
    image     at least one generated picture, downloaded and saved
    video     a rendered file on disk
    data      a reply that the renderer's parser accepts
    links     rows/contacts (Apollo and friends)

Three things follow from a step having a kind, and all three are the point:

1. ONE line in the prompt says what to produce, written for that kind and
   that tool (`deliverable_line`) — replacing the paragraph of "DELIVERABLE
   SPEC / QUALITY BAR / NON-GOALS" the planner used to write for every step
   whether it fitted or not.
2. When the step ends, Prism checks (`missing`). A file step that came back
   as chat text is not "done" — it is asked once for the file
   (`reask`), and if the tool still answers in prose the text itself is
   written out as a real document (`document_from_text`) so the deliverable
   is never lost.
3. The run can say, at the end, which of the things the person asked for
   actually exist — instead of a row of green ticks over an empty folder.

Kinds are deliberately few. Six is enough to cover every stage Prism has,
and a list short enough that a chat model asked to pick one gets it right.
"""
from __future__ import annotations

import os
import re

#: Every kind a step can declare. Anything else is treated as unset.
KINDS = ("text", "file", "image", "video", "data", "links")

#: What a stage produces when nobody said otherwise. The planner may
#: override it per run (a research step asked for a spreadsheet is a `file`
#: step), but a plan that says nothing still knows what each stage is for —
#: which is what makes the contract safe to rely on for old plans, retries
#: and the add-ons that build their own stage lists.
_BY_STAGE = {
    "research": "text",
    "brains": "text",
    "content": "text",
    "summary": "text",
    "analysis": "text",
    "development": "text",
    "leads": "links",
    "visual": "image",
    "artwork": "image",
    "media": "video",
    "audio": "file",
    "presentation": "file",
    "format": "file",      # /boq's own label for its document step
    "write": "file",       # Gerber's
    "script": "text",
    "design": "data",      # read by the renderer, not by a person
    "motion_plan": "data",
}

#: A format named as the thing to RECEIVE. Unambiguous format words match on
#: their own; "excel" and "word" only when a file noun follows, because both
#: are ordinary English ("make them excel at sales", "in a word"). Matched on
#: word edges, never as substrings: the first version matched "excel" inside
#: "excellent" and "pdf" inside "summarise the attached PDF", which turned
#: ordinary writing steps into file steps and wrote Word documents nobody
#: had asked for.
_FMT = (r"(?<![A-Za-z0-9])(?P<fmt>\.?(?:docx|doc|pdf|xlsx|xls|csv|pptx|ppt)\b"
        r"|powerpoint\b|spreadsheet\b"
        r"|(?:excel|word)\s+(?:file|document|doc|sheet|spreadsheet|workbook"
        r"|format)\b)")

#: Request shapes:
#:   "in .docx", "as a PDF", "into a word document"
#:   "DOCX document", "pdf report", "csv export"
#:   "make a pdf", "create an excel sheet of the leads", "send me a docx"
_FILE_REQUEST = [re.compile(p, re.I) for p in (
    r"\b(?:in|as|into)\s+(?:an?\s+)?" + _FMT,
    _FMT + r"\s+(?:file|document|doc|sheet|version|format|copy|report"
           r"|export)\b",
    r"\b(?:make|create|generate|produce|build|prepare|export|give|send"
    r"|convert)\b[^.?!\n]{0,30}?" + _FMT,
)]

#: Asking for a download without naming a format.
_DOWNLOAD_REQUEST = re.compile(
    r"\b(?:downloadable|download link|as a file|send me the file"
    r"|give me the file)\b", re.I)

#: The word right before a format that makes it the INPUT, not the output:
#: "the attached PDF", "this spreadsheet", "your deck".
_INPUT_BEFORE = {"the", "this", "that", "these", "those", "attached",
                 "uploaded", "given", "provided", "your", "my", "our", "his",
                 "her", "their", "its", "existing", "above", "same"}

_EXT_OF = {"docx": ".docx", "doc": ".docx", "word": ".docx", "pdf": ".pdf",
           "xlsx": ".xlsx", "xls": ".xlsx", "excel": ".xlsx",
           "spreadsheet": ".xlsx", "csv": ".csv", "pptx": ".pptx",
           "ppt": ".pptx", "powerpoint": ".pptx"}


def _requested_formats(query: str) -> list[str]:
    """Every format the person asked to RECEIVE, in the order they said it,
    as the key into _EXT_OF."""
    text = query or ""
    found: list[tuple[int, str]] = []
    for pattern in _FILE_REQUEST:
        for m in pattern.finditer(text):
            before = text[:m.start("fmt")].split()
            if before and before[-1].lower().strip(".,:;") in _INPUT_BEFORE:
                continue            # "the attached PDF" is what they gave you
            key = m.group("fmt").split()[0].lower().lstrip(".")
            found.append((m.start("fmt"), key))
    return [key for _pos, key in sorted(found)]


def clean(kind: str) -> str:
    """A kind we recognise, or "" — so a planner typo can never make a step
    unverifiable in a way nothing notices."""
    k = str(kind or "").strip().lower()
    return k if k in KINDS else ""


def wants_file(query: str) -> bool:
    """Did the person ask, in their own words, for something to download?"""
    return bool(_requested_formats(query)
                or _DOWNLOAD_REQUEST.search(query or ""))


def for_stage(stage: str, planner_kind: str = "",
              agent_cfg: dict | None = None) -> str:
    """The kind this step is held to, before the run looks at the request.

    1. the stage's own default — what this kind of step is usually for;
    2. the planner's choice for this run — except that it may not turn a
       picture, video or contacts step into a text step. The planner is
       shown an example with a kind on every line, and a model that copies
       "text" onto "Make the images" would otherwise switch off the image
       line in the prompt, the image check and the end-of-run check at once;
    3. what the tool can physically do (_within_reach).

    The person's own words are applied by file_step_index(), to ONE step —
    not here, where they made every writing step in the run a file step and
    asked a step in the middle for a file and a hand-off at the same time.
    """
    default = _BY_STAGE.get(stage, "text")
    kind = default
    chosen = clean(planner_kind)
    if chosen and not (default in ("image", "video", "links")
                       and chosen == "text"):
        kind = chosen
    return _within_reach(kind, agent_cfg or {})


#: Steps that WRITE, and so can be the one that hands back a requested file.
_WRITING_STAGES = ("content", "brains", "summary", "development", "research",
                   "write", "format")


def file_step_index(entries, query: str) -> int:
    """Which step hands back the file the person asked for, or -1.

    `entries` is [(stage, kind, agent_cfg)] in run order, kinds already
    resolved by for_stage(). One step: the LAST that writes and whose tool
    can return a file — the one whose answer is the finished piece. If the
    plan already has a file step, that step is the answer and nothing is
    promoted. A caller that knows a step is read by a program (a script
    feeding a renderer) passes its kind as "data", which rules it out.
    """
    if not wants_file(query):
        return -1
    if any(kind == "file" for _stage, kind, _cfg in entries):
        return -1
    for i in range(len(entries) - 1, -1, -1):
        stage, kind, cfg = entries[i]
        cfg = cfg or {}
        if kind != "text" or stage not in _WRITING_STAGES or cfg.get("local"):
            continue
        can = cfg.get("produces")
        if can and "file" not in can:
            continue
        return i
    return -1


def _within_reach(kind: str, agent_cfg: dict) -> str:
    """Downgrade a kind the assigned tool cannot deliver.

    `produces` comes from the tool's profile (agents.py). A tool held to
    something it cannot make fails every run with no way for the customer to
    fix it, so it is held to what it CAN make: text when it writes, else the
    first thing it produces — Apollo returns contacts, whatever the plan said.
    """
    can = agent_cfg.get("produces")
    if not can or kind in can:
        return kind
    return "text" if "text" in can else can[0]


# ── what the prompt says ─────────────────────────────────────────────────
#
# ONE line, appended once, in Prism's own voice. Not a section, not a rule
# sheet, and never two of them in the same message: the message that taught
# us this lesson carried the planner's "Deliverable Spec", the engine's
# "HOW THIS FITS IN" and a third block opening "THIS OVERRIDES EVERY OTHER
# FORMATTING INSTRUCTION", and the tool obeyed two of the three at once.

def deliverable_line(kind: str, agent_name: str = "",
                     agent_cfg: dict | None = None, ext: str = "") -> str:
    """The single sentence that says what to hand back.

    `ext` is named only when the person named a format. A deck tool asked
    for "a file" builds a deck; telling it ".docx" — which the first version
    did by default — asks Gamma for a Word document.
    """
    cfg = agent_cfg or {}
    if kind == "file":
        how = (cfg.get("file_hint") or "").strip()
        want = f" ({ext})" if ext else ""
        return ("Give me the finished result as a file I can download"
                f"{want} — build it here in {agent_name or 'this tool'}"
                " rather than pasting it into the chat."
                + (f" {how}" if how else ""))
    if kind == "image":
        return ("Reply with the picture itself, generated here. Not a "
                "description of it, not a prompt for another tool.")
    if kind == "video":
        return ("Reply with the finished video, made here in "
                f"{agent_name or 'this tool'}.")
    if kind == "data":
        return ("Reply in this chat with the JSON object and nothing else — "
                "no preamble, no explanation, and not in a side panel, "
                "canvas or artifact.")
    return ""              # text and links: the brief is the instruction


def label(kind: str) -> str:
    """How the plan screen names this step's deliverable, in plain words."""
    return {
        "file": "a file to download",
        "image": "pictures",
        "video": "a video file",
        "data": "a spec for the renderer",
        "links": "companies and contacts",
        "text": "a written answer",
    }.get(kind, "a written answer")


# ── what Prism checks afterwards ─────────────────────────────────────────

def missing(kind: str, *, texts=None, files: int = 0, images: int = 0) -> str:
    """Why this step did not meet its contract, or "" when it did.

    Deliberately blunt: the string is shown to the person, so it names the
    thing that is not there rather than the mechanism that failed.
    """
    has_text = any((t or "").strip() for t in (texts or []))
    if kind == "file":
        return "" if files else "no file came back — only chat text"
    if kind == "image":
        return "" if images else "no picture was generated"
    if kind == "video":
        return ""              # the renderer reports for itself
    if kind in ("text", "data", "links"):
        return "" if has_text else "nothing came back"
    return ""


def reask(kind: str, ext: str = "") -> str:
    """The one short follow-up sent in the same chat when a step missed its
    contract. Short on purpose — the tool still has everything it wrote a
    moment ago, so this is a correction, not a re-brief."""
    if kind == "file":
        return ("Please give me that as a downloadable "
                f"{ext + ' ' if ext else ''}file, built here — the text in "
                "the chat is not what I need to keep.")
    if kind == "image":
        return ("Please generate the image itself now and show it here — "
                "one picture, no description.")
    if kind == "data":
        return ("Please send that again as the JSON object only, written "
                "in this chat rather than in a side panel.")
    return ""


def wanted_ext(query: str) -> str:
    """The file the person asked for, from their own words. "" when they
    named no format — callers must then not invent one."""
    formats = _requested_formats(query)
    return _EXT_OF.get(formats[0], "") if formats else ""


#: What Prism can write itself from a step's text, if the tool will not.
#: A spreadsheet or a deck built by guessing at the structure of prose would
#: be a worse deliverable than an honest "not produced".
_WRITABLE = ("", ".docx", ".pdf", ".md", ".csv", ".txt")


def can_write(ext: str) -> bool:
    return (ext or "") in _WRITABLE


# ── the last resort: write the document ourselves ────────────────────────

_HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")


def document_from_text(text: str, out_dir: str, stem: str,
                       ext: str = ".docx") -> tuple[str, str]:
    """Write a step's chat answer out as a real document.

    The honest fallback for the case this module was written for: the person
    asked for a document, the tool wrote it in the chat instead, and asking
    once did not change that. The words exist; only the file was missing.
    Writing it here means the run still ends with something in Prism
    Artifacts to open, and the customer is told plainly that Prism made the
    file from the tool's text rather than the tool making it.

    Returns (path, note). ("", "") when there is nothing to write, or the
    format is one Prism cannot build from prose — never an exception, since
    the text itself is already saved in the run.
    """
    text = (text or "").strip()
    if not text or not can_write(ext):
        return "", ""
    os.makedirs(out_dir, exist_ok=True)
    # A PDF is the one common ask Prism cannot honour on its own — no PDF
    # renderer is bundled — so the Word file is written instead and the
    # difference is said out loud rather than left for the customer to find.
    if ext in ("", ".docx", ".pdf"):
        try:
            path = _write_docx(text, os.path.join(out_dir, stem + ".docx"))
            note = "Prism saved the step's own text as a Word document"
            if ext == ".pdf":
                note += " (a PDF can't be made on this computer)"
            return path, note
        except Exception:                                    # noqa: BLE001
            ext = ".md"        # python-docx missing or unhappy — fall through

    try:
        path = os.path.join(out_dir, stem + ext)
        with open(path, "w", encoding="utf-8") as f:
            f.write(text + "\n")
        return path, "Prism saved the step's own text as a file"
    except OSError:
        return "", ""


def _write_docx(text: str, path: str) -> str:
    """Markdown-ish chat text into a Word file with real headings.

    Not a Markdown renderer — deliberately. What arrives is what a chat tool
    writes: '## Section', '- bullet', blank lines between paragraphs. Those
    three shapes carry the structure a reader needs; anything cleverer would
    be guessing at the rest.
    """
    from docx import Document                       # noqa: PLC0415
    doc = Document()
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        head = _HEADING.match(line)
        if head:
            doc.add_heading(head.group(2).strip(), level=min(len(head.group(1)), 4))
            continue
        bullet = _BULLET.match(line)
        if bullet:
            doc.add_paragraph(bullet.group(1).strip(), style="List Bullet")
            continue
        doc.add_paragraph(line.strip())
    doc.save(path)
    return path


# ── what the whole run owed the person ───────────────────────────────────

def promised(stages) -> list[str]:
    """The kinds this run is expected to end with, from its resolved steps.

    `stages` is [(stage, kind)] AFTER file_step_index() has been applied, so
    a requested file is already carried by the step that owes it. Kept here
    rather than in the run loop so the CLI, the GUI and the tests all answer
    the question the same way.
    """
    want: list[str] = []
    for _stage, kind in stages:
        if kind in ("file", "image", "video") and kind not in want:
            want.append(kind)
    return want


def shortfall(want: list[str], got: dict) -> list[str]:
    """Which promised deliverables are not on disk. `got` maps kind -> count."""
    return [k for k in want if not got.get(k)]
