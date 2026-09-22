"""
Prism — file attachments (any type)
───────────────────────────────────
Users can attach ANY file to a task. Prism does two things with each one:

  1. Extracts text when possible (txt/md/code/csv/json/pdf/docx/…) and injects
     it into the routing brain + every agent's context, so even tools without an
     upload box still "see" the content.
  2. Uploads the real file to each tool's web UI (via its <input type="file">),
     so images, audio, video, zips, datasets, etc. reach agents that accept them.

Text extraction for PDF/DOCX uses optional deps (pypdf, python-docx). If they're
missing, the file is still attached and uploaded — only inline text is skipped.
"""
from __future__ import annotations
import os
import mimetypes

# How much extracted text to inline per file (keeps prompts sane).
MAX_TEXT_CHARS = 12000
# When the same file has ALSO been uploaded to the tool, inline only this
# much: past it the tool reads the attachment, and the prompt carries a
# description instead. A 370 KB attendance sheet pasted into the prompt as
# 12,000 characters of truncated rows — beside the very same file as an
# upload — is what stalled the analysis stage on a real run.
INLINE_WHEN_UPLOADED = 2500
# Spreadsheet-like files (csv/tsv/xlsx) above this many rows are never
# inlined raw; they get a profile — columns, row count, per-column summary,
# a few sample rows — which is what "analyse this file" needs anyway.
PROFILE_ABOVE_ROWS = 40
_TABULAR_EXTS = {".csv", ".tsv", ".xlsx"}

# Extensions we treat as directly-readable UTF-8 text even when mimetypes is unsure.
_TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json", ".jsonl",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".env", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".h", ".cpp", ".cc", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".kt", ".scala", ".sh", ".bash",
    ".zsh", ".sql", ".r", ".m", ".lua", ".pl", ".dart", ".vue", ".svelte", ".css",
    ".scss", ".less", ".tex", ".bib", ".srt", ".vtt",
}


def _human_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f}{unit}" if unit == "B" else f"{n:.1f}{unit}"
        n /= 1024
    return f"{n:.1f}GB"


def _classify(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    mime, _ = mimetypes.guess_type(path)
    if ext in _TEXT_EXTS:
        return "text"
    if mime:
        if mime.startswith("image/"):
            return "image"
        if mime.startswith("audio/"):
            return "audio"
        if mime.startswith("video/"):
            return "video"
        if mime.startswith("text/"):
            return "text"
        if mime in ("application/json", "application/xml",
                    "application/x-yaml", "application/javascript"):
            return "text"
    if ext == ".pdf":
        return "pdf"
    if ext in (".docx",):
        return "docx"
    return "binary"


def _docx_text(document) -> str | None:
    """Everything a .docx says — paragraphs AND tables, in document order.

    This was `"\\n".join(p.text for p in d.paragraphs)`, and `d.paragraphs`
    is body paragraphs ONLY: it does not descend into tables. Measured on
    this repo's own sample_boq.docx, that read 1,924 characters and dropped
    11,214 — nine tables, 111 rows, 85% of the file. A bill of quantities, a
    quotation, a rate list and a purchase order are all *mostly table*, so
    the part being thrown away was the part the customer attached the file
    for; the tool then answered from a heading and an intro paragraph.

    Order matters as much as inclusion. Appending the tables after all the
    prose would put a table's numbers under the wrong heading, which is a
    worse kind of wrong than leaving them out — so the body's own children
    are walked in sequence instead.

    Matched on tag name rather than by importing docx.oxml's CT_P/CT_Tbl:
    those are internal classes that have moved between python-docx releases,
    and this file is read by a frozen build that cannot be patched quickly.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(child, document).text)
        elif tag == "tbl":
            try:
                for row in Table(child, document).rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
            except Exception:       # noqa: BLE001
                continue            # a malformed table must not lose the prose
    text = "\n".join(lines).strip()
    return text or None


def _extract_text(path: str, kind: str) -> str | None:
    try:
        if kind == "text":
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(MAX_TEXT_CHARS + 1)
        if kind == "pdf":
            try:
                from pypdf import PdfReader
            except Exception:
                try:
                    from PyPDF2 import PdfReader  # older name
                except Exception:
                    return None
            reader = PdfReader(path)
            out = []
            for page in reader.pages:
                out.append(page.extract_text() or "")
                if sum(len(x) for x in out) > MAX_TEXT_CHARS:
                    break
            return "\n".join(out) or None
        if kind == "docx":
            try:
                import docx  # python-docx
            except Exception:
                return None
            return _docx_text(docx.Document(path))
        # Last resort: sniff whether an unknown binary is actually decodable text.
        with open(path, "rb") as f:
            head = f.read(4096)
        if b"\x00" not in head:
            try:
                text = head.decode("utf-8")
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read(MAX_TEXT_CHARS + 1)
            except Exception:
                return None
    except Exception:
        return None
    return None


# Folder attach cap — ChatGPT (the file-analysis stage) takes at most 20
# files per message, so stay comfortably under it.
MAX_DIR_FILES = 15


def attach_dir(path: str) -> list[dict]:
    """Attachment records for every plain file directly inside a folder,
    falling back to subdirectories if none are in the root (hidden files
    skipped, capped at MAX_DIR_FILES)."""
    path = os.path.abspath(os.path.expanduser(path))
    if not os.path.isdir(path):
        raise NotADirectoryError(path)
    out = []
    for name in sorted(os.listdir(path)):
        fp = os.path.join(path, name)
        if name.startswith(".") or not os.path.isfile(fp):
            continue
        try:
            out.append(attach(fp))
        except Exception:
            continue
        if len(out) >= MAX_DIR_FILES:
            return out
    if not out:
        for root, dirs, files in os.walk(path):
            dirs[:] = [d for d in dirs if not d.startswith(".")]
            for name in sorted(files):
                if name.startswith("."):
                    continue
                fp = os.path.join(root, name)
                try:
                    out.append(attach(fp))
                except Exception:
                    continue
                if len(out) >= MAX_DIR_FILES:
                    return out
    return out


def attach(path: str) -> dict:
    """Build an attachment record for any file. Raises if it doesn't exist."""
    path = os.path.abspath(os.path.expanduser(path))
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    if os.path.isdir(path):
        raise IsADirectoryError(path)
    kind = _classify(path)
    text = _extract_text(path, kind)
    truncated = bool(text and len(text) > MAX_TEXT_CHARS)
    if truncated:
        text = text[:MAX_TEXT_CHARS]
    profile = ""
    if os.path.splitext(path)[1].lower() in _TABULAR_EXTS:
        profile = tabular_profile(path)
    return {
        "path": path,
        "name": os.path.basename(path),
        "size": os.path.getsize(path),
        "mime": mimetypes.guess_type(path)[0] or "application/octet-stream",
        "kind": kind,
        "text": text,
        "truncated": truncated,
        # For a spreadsheet-like file: what it holds, in a few hundred
        # characters, instead of a wall of rows. Empty for everything else.
        "profile": profile,
    }


# ── spreadsheets: a profile, not a wall of rows ──────────────────────────────

def _read_table(path: str, limit: int = 200_000) -> tuple[list[str], list[list[str]]]:
    """Header and rows of a csv/tsv/xlsx, capped so a huge file stays cheap."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        try:
            import openpyxl
        except Exception:
            return [], []
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v) for v in row])
            if len(rows) > limit:
                break
        wb.close()
    else:
        import csv
        delimiter = "\t" if ext == ".tsv" else ","
        with open(path, "r", encoding="utf-8-sig", errors="ignore", newline="") as f:
            rows = []
            for row in csv.reader(f, delimiter=delimiter):
                rows.append(row)
                if len(rows) > limit:
                    break
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return [], []
    return [c.strip() for c in rows[0]], rows[1:]


def _as_number(text: str):
    t = text.strip().replace(",", "")
    if not t:
        return None
    try:
        return float(t)
    except ValueError:
        return None


def tabular_profile(path: str, sample_rows: int = 6, max_chars: int = 3500) -> str:
    """What a spreadsheet holds, in a few hundred characters.

    Row count, the columns, and for each column the kind of thing in it:
    a range for numbers and dates, the distinct values with counts for a
    short list (Status: Present 6,912 · Absent 401 · Leave 187), the count
    of distinct values with examples for everything else — then a handful
    of rows verbatim so the reader can see the shape. This is the brief an
    "analyse this file" stage would spend its whole turn producing from the
    raw rows, and it is exact where a model's own count would be a guess.

    Never raises; an unreadable file profiles as an empty string.
    """
    try:
        header, rows = _read_table(path)
    except Exception:
        return ""
    if not header:
        return ""
    lines = [f"{len(rows):,} rows x {len(header)} columns. Columns: "
             + ", ".join(header)]
    for i, name in enumerate(header):
        values = [r[i].strip() for r in rows if i < len(r)]
        filled = [v for v in values if v]
        if not filled:
            lines.append(f"  {name}: empty")
            continue
        blanks = len(values) - len(filled)
        blank_note = f"; {blanks:,} blank" if blanks else ""
        numbers = [n for n in (_as_number(v) for v in filled) if n is not None]
        distinct = {}
        for v in filled:
            distinct[v] = distinct.get(v, 0) + 1
        if len(numbers) == len(filled) and len(distinct) > 12:
            lo, hi = min(numbers), max(numbers)
            mean = sum(numbers) / len(numbers)
            fmt = (lambda x: f"{x:,.0f}") if all(n == int(n) for n in numbers) \
                else (lambda x: f"{x:,.2f}")
            lines.append(f"  {name}: numbers {fmt(lo)} to {fmt(hi)}, "
                         f"average {fmt(mean)}{blank_note}")
        elif len(distinct) <= 12:
            top = sorted(distinct.items(), key=lambda kv: -kv[1])
            lines.append(f"  {name}: " + " · ".join(f"{v} {n:,}" for v, n in top)
                         + blank_note)
        else:
            import re as _re
            if all(_re.fullmatch(r"\d{1,2}:\d{2}", v) for v in list(distinct)[:200]):
                # Clock times / durations: order by minutes, not as text,
                # or "from 10:00 to 9:59" comes out backwards.
                ordered = sorted(distinct, key=lambda v: int(v.split(":")[0]) * 60
                                 + int(v.split(":")[1]))
                looks_sorted = True
            else:
                ordered = sorted(distinct)
                looks_sorted = all(len(v) == len(ordered[0]) for v in ordered[:50])
            if looks_sorted and (ordered[0][:2].isdigit() or ":" in ordered[0]):
                lines.append(f"  {name}: {len(distinct):,} distinct, from "
                             f"{ordered[0]} to {ordered[-1]}{blank_note}")
            else:
                examples = ", ".join(list(distinct)[:3])
                lines.append(f"  {name}: {len(distinct):,} distinct, e.g. "
                             f"{examples}{blank_note}")
    lines.append(f"First {min(sample_rows, len(rows))} rows, as written:")
    lines.append("  " + " | ".join(header))
    for r in rows[:sample_rows]:
        lines.append("  " + " | ".join(r))
    out = "\n".join(lines)
    return out[:max_chars]


def describe(att: dict) -> str:
    tag = "📄 text" if att["text"] else f"📎 {att['kind']}"
    return f"{att['name']}  [dim]({_human_size(att['size'])} · {tag})[/dim]"


def _row_count(att: dict) -> int:
    head = (att.get("profile") or "").split(" rows", 1)[0].replace(",", "")
    return int(head) if head.isdigit() else 0


def context_block(attachments: list[dict], uploaded: bool = False) -> str:
    """Text injected into agent prompts so tools see file contents inline.

    `uploaded` — the same files have just gone up to this tool as real
    attachments. Then a big text file is NOT pasted in as well: the prompt
    says it is attached and, for a spreadsheet, gives the profile. Pasting
    it too is how a 370 KB sheet became 12,000 characters of truncated rows
    beside its own upload, and the tool spent the turn reading them.
    """
    if not attachments:
        return ""
    parts = ["📎 The user attached the following file(s). Use them as primary source material:\n"]
    for att in attachments:
        header = f"── {att['name']} ({_human_size(att['size'])}, {att['kind']})"
        profile = att.get("profile") or ""
        text = att.get("text") or ""
        big_table = bool(profile) and _row_count(att) > PROFILE_ABOVE_ROWS
        if big_table:
            where = ("uploaded to this chat as a file — read the rows there"
                     if uploaded else
                     "too large to paste; only this profile is given")
            parts.append(f"{header}  [spreadsheet: {where}]\n{profile}\n")
        elif text and uploaded and len(text) > INLINE_WHEN_UPLOADED:
            parts.append(f"{header}\n(uploaded to this chat as a file — read it "
                         "from the attachment; not pasted here)\n")
        elif text:
            trunc = "  [content truncated]" if att.get("truncated") else ""
            parts.append(f"{header}{trunc}\n{text}\n")
        else:
            parts.append(f"{header}\n(binary file — uploaded directly to the tool; contents not inlined)\n")
    parts.append("── end of attachments ──\n")
    return "\n".join(parts) + "\n"


def routing_note(attachments: list[dict]) -> str:
    """A short note for the routing brain (full contents are given to agents)."""
    if not attachments:
        return ""
    names = ", ".join(a["name"] for a in attachments)
    kinds = ", ".join(sorted({a["kind"] for a in attachments}))
    return (
        f"\nThe user attached {len(attachments)} file(s): {names} (types: {kinds}). "
        "Their contents are provided to each agent as context and the raw files are "
        "uploaded to each tool. Write prompts that explicitly use these files.\n"
    )


def upload_paths(attachments: list[dict]) -> list[str]:
    return [a["path"] for a in attachments]
